import logging
from pathlib import Path
from typing import List, Optional, Union
import docx

from src.models.document import DocumentPage, ExtractionMethod
from src.parsers.base import Parser
from src.parsers.table_extractor import TableExtractor

logger = logging.getLogger("medical_rag.parsers.docx_parser")

MIN_MEANINGFUL_CHAR_COUNT = 30


class DOCXParser(Parser):
    """DOCX Document Parser using python-docx with heading structure preservation and table extraction."""

    def __init__(self, table_extractor: Optional[TableExtractor] = None):
        self.table_extractor = table_extractor or TableExtractor()

    def parse(self, file_path: Union[str, Path], document_id: str) -> List[DocumentPage]:
        """Parses a DOCX document into a single DocumentPage instance.

        Args:
            file_path: Path to the target DOCX file.
            document_id: UUID of parent Document entity.

        Returns:
            List containing 1 DocumentPage.
        """
        path_str = str(file_path)
        if not Path(path_str).exists():
            raise FileNotFoundError(f"DOCX file not found at path: '{path_str}'")

        try:
            logger.info(f"Opening DOCX file '{path_str}' for document_id='{document_id}'")
            doc = docx.Document(path_str)

            paragraph_lines = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                # Format headings with Markdown hashes if heading style detected
                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    heading_level = 1
                    try:
                        # Extract heading level digit if present (e.g., 'Heading 2' -> 2)
                        heading_level = int(p.style.name.split()[-1])
                    except (ValueError, IndexError):
                        heading_level = 1
                    prefix = "#" * max(1, min(6, heading_level))
                    paragraph_lines.append(f"{prefix} {text}")
                else:
                    paragraph_lines.append(text)

            raw_text = "\n\n".join(paragraph_lines)
            stripped_text = raw_text.strip()

            has_meaningful_text = len(stripped_text) >= MIN_MEANINGFUL_CHAR_COUNT
            extraction_method = (
                ExtractionMethod.NATIVE if has_meaningful_text else ExtractionMethod.OCR
            )

            tables = self.table_extractor.extract_from_docx_tables(doc.tables)

            page = DocumentPage(
                document_id=document_id,
                page_number=1,
                raw_text=raw_text,
                cleaned_text=raw_text,
                extraction_method=extraction_method,
                tables=tables,
                has_meaningful_text=has_meaningful_text
            )

            logger.info(f"Parsed DOCX '{path_str}': single page extracted ({len(raw_text)} chars)")
            return [page]
        except Exception as err:
            logger.error(f"Failed to parse DOCX file '{path_str}': {err}")
            raise RuntimeError(f"Failed to parse DOCX document: {str(err)}") from err
